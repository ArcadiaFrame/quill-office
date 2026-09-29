#!/usr/bin/env python3
"""Build the original Task Launcher templates (documents, spreadsheets, presentations).

All content is written for this project (no third-party templates). Output goes to
aero/templates/out/{Documents,Spreadsheet,Presentation}/<Name>.<ext>. apply-aero.py copies
it into <app>/converter/templates/EN/..., where the app lists the files in its Templates
panel and renders their thumbnails. The launcher refers to them as "builtin:<Name>".

Needs: pip install python-docx openpyxl python-pptx
Usage: python aero/templates/build_templates.py
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Cm

OUT = Path(__file__).resolve().parent / "out"
NAVY = RGBColor(0x1D, 0x35, 0x57)
BLUE = RGBColor(0x2F, 0x6F, 0xB5)
GREY = RGBColor(0x5F, 0x6B, 0x7A)
NAVY_HEX, BLUE_HEX, PALE_HEX = "1D3557", "2F6FB5", "E8F0F9"


# ── Word helpers ─────────────────────────────────────────────────────────────
def new_doc(body_font="Lato", size=11, margins_in=1.0):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = body_font
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
    st.paragraph_format.space_after = Pt(6)
    for name, sz, color in (("Heading 1", 18, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 12, NAVY), ("Title", 30, NAVY)):
        h = d.styles[name]
        h.font.name = "Montserrat"
        h.font.size = Pt(sz)
        h.font.color.rgb = color
        h.font.bold = True
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Montserrat")
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(margins_in)
    return d


def para(d, text="", size=None, bold=False, italic=False, color=None, align=None, font=None, after=None, before=None, style=None):
    p = d.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
        if font:
            r.font.name = font
            r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    return p


def rule(d, color=BLUE_HEX, size=12):
    """A colored bottom border on an empty paragraph (a horizontal line)."""
    p = d.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "1"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)
    p.paragraph_format.space_after = Pt(8)
    return p


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(d, rows, header=True, widths=None, header_fill=NAVY_HEX):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            if header and i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, header_fill)
            elif i % 2 == 0:
                shade(c, PALE_HEX)
            if widths: c.width = widths[j]
    d.add_paragraph()
    return t


def columns(section, n=2, space_twips=432):
    cols = section._sectPr.xpath("./w:cols")
    col = cols[0] if cols else OxmlElement("w:cols")
    col.set(qn("w:num"), str(n)); col.set(qn("w:space"), str(space_twips))
    if not cols: section._sectPr.append(col)


def core(d, title):
    d.core_properties.title = title
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""


def save(d, folder, name):
    path = OUT / folder / (name + ".docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    core(d, name)
    d.save(path)
    return path


# ── Documents ─────────────────────────────────────────────────────────────────
def business_letter():
    d = new_doc()
    para(d, "[Your Company]", 20, bold=True, color=NAVY, font="Montserrat", after=0)
    para(d, "[Street Address] · [City, State ZIP] · [Phone] · [Email]", 9, color=GREY, after=0)
    rule(d)
    para(d, "[Month Day, Year]", after=14)
    for line in ("[Recipient Name]", "[Title]", "[Company]", "[Street Address]", "[City, State ZIP]"):
        para(d, line, after=0)
    para(d, "")
    para(d, "Re: [Subject of the letter]", bold=True, after=12)
    para(d, "Dear [Recipient Name],", after=10)
    para(d, "[Opening: state the purpose of your letter in one or two sentences.]")
    para(d, "[Body: give the details, background or request. Keep paragraphs short and focused on one idea each.]")
    para(d, "[Closing: summarize what happens next or what you would like the reader to do.]")
    para(d, "Sincerely,", before=12, after=36)
    para(d, "[Your Name]", bold=True, after=0)
    para(d, "[Your Title]", color=GREY)
    return save(d, "Documents", "Business letter")


def cover_letter():
    d = new_doc()
    para(d, "[YOUR NAME]", 24, bold=True, color=NAVY, font="Montserrat", after=0)
    para(d, "[Phone] · [Email] · [City, State] · [LinkedIn or portfolio]", 9, color=GREY, after=0)
    rule(d)
    para(d, "[Month Day, Year]", after=10)
    for line in ("[Hiring Manager Name]", "[Company]", "[Address]"):
        para(d, line, after=0)
    para(d, "")
    para(d, "Dear [Hiring Manager Name],", after=10)
    para(d, "I am excited to apply for the [Job Title] position at [Company]. [One sentence on why this role and company appeal to you.]")
    para(d, "In my role as [Current or Recent Role] at [Organization], I [key accomplishment with a number, e.g. increased X by 20%]. "
            "I also [second relevant achievement], which prepared me to [something the job requires].")
    para(d, "What draws me to [Company] is [specific detail about the company, product or mission]. I would bring [skill] and [skill] to your team.")
    para(d, "Thank you for your time and consideration. I would welcome the chance to discuss how I can contribute.")
    para(d, "Sincerely,", before=12, after=30)
    para(d, "[Your Name]", bold=True)
    return save(d, "Documents", "Cover letter")


def thank_you_note():
    d = new_doc(body_font="Noto Serif", size=12, margins_in=1.25)
    para(d, "")
    para(d, "Thank you", 48, color=NAVY, font="Caveat", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    rule(d, BLUE_HEX, 6)
    para(d, "Dear [Name],", before=18, after=12)
    para(d, "Thank you so much for [the gift, your help, coming to …]. [Say what it meant to you or how you will use or remember it.]", after=12)
    para(d, "[Add a personal line: a shared memory, a hope to see them soon, or a wish for them.]", after=18)
    para(d, "With gratitude,", after=6)
    para(d, "[Your Name]", 22, font="Caveat", color=BLUE)
    return save(d, "Documents", "Thank-you note")


def event_flyer():
    d = new_doc(margins_in=0.6)
    t = d.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, NAVY_HEX)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("YOU'RE INVITED")
    r.font.size, r.font.bold, r.font.color.rgb = Pt(14), True, RGBColor(0xBF, 0xD7, 0xF2)
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("[EVENT NAME]")
    r2.font.name = "Bebas Neue"; r2.font.size = Pt(72); r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p3 = c.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("[A one-line description that makes people want to come]")
    r3.font.size, r3.font.color.rgb = Pt(14), RGBColor(0xE8, 0xF0, 0xF9)
    para(d, "")
    table(d, [["WHEN", "WHERE", "COST"], ["[Day, Month Date]\n[Start – End time]", "[Venue name]\n[Street address]", "[Free / $X per person]"]])
    para(d, "What to expect", 18, bold=True, color=NAVY, font="Montserrat", align=WD_ALIGN_PARAGRAPH.CENTER)
    for item in ("[Highlight one: music, speakers, food…]", "[Highlight two]", "[Highlight three]"):
        para(d, "•  " + item, 13, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(d, "RSVP by [date] · [phone or email] · [website]", 12, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=18)
    return save(d, "Documents", "Event flyer")


def newsletter():
    d = new_doc(size=10, margins_in=0.7)
    para(d, "[ORGANIZATION] NEWS", 32, bold=True, color=NAVY, font="Montserrat", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(d, "[Month Year] · Issue [#] · [Tagline]", 10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    rule(d)
    sec = d.add_section(0)
    sec.left_margin = sec.right_margin = Inches(0.7)
    columns(sec, 2)
    d.add_heading("[Lead story headline]", level=1)
    para(d, "[Open with the most important news of this issue. Who, what, when, where and why, in two or three short paragraphs.]")
    para(d, "[Add a quote from someone involved to bring the story to life.]", italic=True, color=GREY)
    d.add_heading("[Second story]", level=2)
    para(d, "[A shorter update: a new project, a milestone or a recognition.]")
    d.add_heading("Upcoming dates", level=2)
    for item in ("[Date] – [Event]", "[Date] – [Event]", "[Date] – [Event]"):
        para(d, "•  " + item, after=2)
    d.add_heading("Spotlight", level=2)
    para(d, "[Introduce a person, team or volunteer. What do they do, and what is something surprising about them?]")
    d.add_heading("Get in touch", level=2)
    para(d, "[Contact name] · [email] · [phone] · [website]")
    return save(d, "Documents", "Newsletter")


def research_report():
    d = new_doc(body_font="Noto Serif", size=11)
    for _ in range(6): para(d, "")
    para(d, "[Report Title]", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d, "[Subtitle or research question]", 14, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
    para(d, "[Your Name]", 12, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(d, "[Course or Organization]", 12, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(d, "[Date]", 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for h, lvl, body in (("Abstract", 1, "[A 150–250 word summary: the question, the method, the main findings and why they matter.]"),
                         ("Introduction", 1, "[Background on the topic and why it matters. End with your research question or thesis.]"),
                         ("Method", 1, "[How you gathered and analysed information: sources, surveys, experiments or data sets.]"),
                         ("Results", 1, "[What you found. Use tables or charts, and describe them in words too.]"),
                         ("Key finding one", 2, "[Detail and evidence.]"),
                         ("Key finding two", 2, "[Detail and evidence.]"),
                         ("Discussion", 1, "[What the results mean, how they compare with other work, and their limits.]"),
                         ("Conclusion", 1, "[Answer the research question and suggest next steps.]"),
                         ("References", 1, "[Author, A. (Year). Title of work. Publisher or Journal, volume(issue), pages.]")):
        d.add_heading(h, level=lvl)
        para(d, body)
    return save(d, "Documents", "Research report")


def book_report():
    d = new_doc(size=12)
    para(d, "Book Report", style="Title")
    table(d, [["Title", "[Book title]"], ["Author", "[Author name]"], ["Genre", "[Genre]"], ["Pages", "[Number]"],
              ["Student", "[Your name]"], ["Date", "[Date]"]], header=False)
    for h, body in (("Summary", "[In your own words, what happens in the book? Cover the beginning, middle and end without retelling every detail.]"),
                    ("Main characters", "[Who are the important characters? Describe each in one or two sentences.]"),
                    ("Setting", "[Where and when does the story take place? How does the setting affect the story?]"),
                    ("Themes", "[What big ideas does the book explore, such as friendship, courage or change?]"),
                    ("My opinion", "[Did you enjoy it? Why or why not? Who would you recommend it to?]"),
                    ("Favourite quote", "“[Quote]” – page [#]")):
        d.add_heading(h, level=1)
        para(d, body)
    para(d, "Rating:  ☆ ☆ ☆ ☆ ☆", 14, bold=True, color=NAVY, before=10)
    return save(d, "Documents", "Book report")


def meeting_minutes():
    d = new_doc()
    para(d, "Meeting Minutes", style="Title")
    table(d, [["Meeting", "[Name or purpose]"], ["Date and time", "[Date], [Start] – [End]"], ["Location", "[Room / video link]"],
              ["Facilitator", "[Name]"], ["Note taker", "[Name]"]], header=False)
    d.add_heading("Attendees", level=1)
    para(d, "[Name], [Name], [Name] · Absent: [Name]")
    d.add_heading("Agenda and discussion", level=1)
    for i in range(1, 4):
        d.add_heading(f"{i}. [Agenda item]", level=2)
        para(d, "[Key points discussed and any decisions made.]")
    d.add_heading("Action items", level=1)
    table(d, [["Action", "Owner", "Due date", "Status"], ["[Task]", "[Name]", "[Date]", "Open"],
              ["[Task]", "[Name]", "[Date]", "Open"], ["[Task]", "[Name]", "[Date]", "Open"]])
    d.add_heading("Next meeting", level=1)
    para(d, "[Date, time and location]")
    return save(d, "Documents", "Meeting minutes")


def memo():
    d = new_doc()
    para(d, "MEMO", 36, bold=True, color=NAVY, font="Montserrat", after=6)
    for k, v in (("To:", "[Recipients]"), ("From:", "[Your name, title]"), ("CC:", "[Others]"), ("Date:", "[Date]"), ("Subject:", "[Topic]")):
        p = d.add_paragraph()
        r = p.add_run(k.ljust(10)); r.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(2)
    rule(d)
    para(d, "[State the purpose of the memo in the first sentence.]")
    para(d, "[Give the context and the details the readers need. Use a short list if there are several points:]")
    for item in ("[Point one]", "[Point two]", "[Point three]"):
        para(d, "•  " + item, after=2)
    para(d, "[End with the action you need from the readers and any deadline.]", before=8)
    return save(d, "Documents", "Memo")


def recipe_card():
    d = new_doc(size=11, margins_in=0.9)
    para(d, "[Recipe Name]", 34, color=NAVY, font="Caveat", after=0)
    para(d, "Serves [#] · Prep [#] min · Cook [#] min", 10, color=GREY)
    rule(d)
    d.add_heading("Ingredients", level=2)
    for item in ("[Amount] [ingredient]", "[Amount] [ingredient]", "[Amount] [ingredient]", "[Amount] [ingredient]", "[Amount] [ingredient]"):
        para(d, "☐  " + item, after=2)
    d.add_heading("Directions", level=2)
    for i in range(1, 5):
        para(d, f"{i}.  [Step]", after=4)
    d.add_heading("Notes", level=2)
    para(d, "[Substitutions, storage tips or where the recipe came from.]", italic=True, color=GREY)
    return save(d, "Documents", "Recipe card")


DOCUMENTS = [business_letter, cover_letter, thank_you_note, event_flyer, newsletter, research_report,
             book_report, meeting_minutes, memo, recipe_card]


# ── Spreadsheets ──────────────────────────────────────────────────────────────
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor=NAVY_HEX)
PALE_FILL = PatternFill("solid", fgColor=PALE_HEX)
THIN = Side(style="thin", color="B7C7DA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet_title(ws, title, sub, width_cols):
    ws["A1"] = title
    ws["A1"].font = Font(name="Montserrat", size=20, bold=True, color=NAVY_HEX)
    ws["A2"] = sub
    ws["A2"].font = Font(name="Lato", size=10, color="5F6B7A")
    for c in range(1, width_cols + 1):
        ws.cell(row=3, column=c).border = Border(bottom=Side(style="medium", color=BLUE_HEX))


def header_row(ws, row, values, start_col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font = Font(name="Lato", bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX


def body(ws, r1, r2, c1, c2, stripe=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            if cell.font is None or not cell.font.bold:
                cell.font = Font(name="Lato", size=11)
            if stripe and (r - r1) % 2 == 1:
                cell.fill = PALE_FILL


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def save_wb(wb, name):
    path = OUT / "Spreadsheet" / (name + ".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.properties.title = name
    wb.properties.creator = ""
    wb.save(path)
    return path


def monthly_budget():
    wb = Workbook(); ws = wb.active; ws.title = "Budget"
    sheet_title(ws, "Monthly Budget", "[Month Year] · Enter planned and actual amounts; totals update automatically.", 4)
    header_row(ws, 5, ["Income", "Planned", "Actual", "Difference"])
    inc = ["Salary / wages", "Side income", "Other"]
    for i, n in enumerate(inc):
        r = 6 + i
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=4, value=f"=C{r}-B{r}")
    tr = 6 + len(inc)
    body(ws, 6, tr, 1, 4)
    ws.cell(row=tr, column=1, value="Total income").font = Font(name="Lato", bold=True)
    for col in "BCD":
        ws[f"{col}{tr}"] = f"=SUM({col}6:{col}{tr - 1})"
    er0 = tr + 2
    header_row(ws, er0, ["Expenses", "Planned", "Actual", "Difference"])
    exp = ["Rent / mortgage", "Utilities", "Groceries", "Transportation", "Insurance", "Phone & internet",
           "Health", "Debt payments", "Savings", "Entertainment", "Other"]
    for i, n in enumerate(exp):
        r = er0 + 1 + i
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
    et = er0 + 1 + len(exp)
    body(ws, er0 + 1, et, 1, 4)
    ws.cell(row=et, column=1, value="Total expenses").font = Font(name="Lato", bold=True)
    for col in "BCD":
        ws[f"{col}{et}"] = f"=SUM({col}{er0 + 1}:{col}{et - 1})"
    net = et + 2
    ws.cell(row=net, column=1, value="Left over (income − expenses)").font = Font(name="Lato", bold=True, color=NAVY_HEX, size=12)
    ws[f"B{net}"] = f"=B{tr}-B{et}"
    ws[f"C{net}"] = f"=C{tr}-C{et}"
    for col in "BC":
        ws[f"{col}{net}"].font = Font(name="Lato", bold=True, size=12, color=NAVY_HEX)
    for r in range(6, net + 1):
        for col in "BCD":
            ws[f"{col}{r}"].number_format = '#,##0.00'
    widths(ws, [30, 14, 14, 14])
    return save_wb(wb, "Monthly budget")


def study_schedule():
    wb = Workbook(); ws = wb.active; ws.title = "Schedule"
    sheet_title(ws, "Weekly Study Schedule", "Fill each block with a subject or task. Colour-code subjects if you like.", 8)
    header_row(ws, 5, ["Time", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"])
    times = [f"{h}:00" for h in range(7, 22)]
    for i, t in enumerate(times):
        ws.cell(row=6 + i, column=1, value=t).font = Font(name="Lato", bold=True)
    body(ws, 6, 5 + len(times), 1, 8)
    ws.cell(row=7 + len(times), column=1, value="Goals this week:").font = Font(name="Lato", bold=True, color=NAVY_HEX)
    widths(ws, [9] + [16] * 7)
    ws.freeze_panes = "B6"
    return save_wb(wb, "Study schedule")


def grade_tracker():
    wb = Workbook(); ws = wb.active; ws.title = "Grades"
    sheet_title(ws, "Grade Tracker", "[Course] · Enter scores and weights; the weighted average updates automatically.", 6)
    header_row(ws, 5, ["Assignment", "Type", "Score", "Out of", "Percent", "Weight %"])
    items = [("[Homework 1]", "Homework"), ("[Quiz 1]", "Quiz"), ("[Project]", "Project"), ("[Midterm]", "Exam"),
             ("[Homework 2]", "Homework"), ("[Quiz 2]", "Quiz"), ("[Final]", "Exam")]
    for i, (n, t) in enumerate(items):
        r = 6 + i
        ws.cell(row=r, column=1, value=n); ws.cell(row=r, column=2, value=t)
        ws.cell(row=r, column=5, value=f'=IF(D{r}>0,C{r}/D{r},"")').number_format = "0.0%"
    last = 5 + len(items)
    body(ws, 6, last, 1, 6)
    ws.cell(row=last + 2, column=4, value="Weighted average").font = Font(name="Lato", bold=True, color=NAVY_HEX)
    c = ws.cell(row=last + 2, column=5, value=f'=IFERROR(SUMPRODUCT(E6:E{last},F6:F{last})/SUM(F6:F{last}),"")')
    c.number_format = "0.0%"
    c.font = Font(name="Lato", bold=True, size=12, color=NAVY_HEX)
    widths(ws, [24, 14, 10, 10, 12, 12])
    return save_wb(wb, "Grade tracker")


def chore_chart():
    wb = Workbook(); ws = wb.active; ws.title = "Chores"
    sheet_title(ws, "Weekly Chore Chart", "Write a name in each box, and tick it off when the chore is done.", 8)
    header_row(ws, 5, ["Chore", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"])
    chores = ["Make beds", "Dishes", "Laundry", "Vacuum", "Take out trash", "Water plants", "Feed pets", "Tidy living room"]
    for i, ch in enumerate(chores):
        ws.cell(row=6 + i, column=1, value=ch).font = Font(name="Lato", bold=True)
    body(ws, 6, 5 + len(chores), 1, 8)
    widths(ws, [22] + [11] * 7)
    return save_wb(wb, "Chore chart")


def meal_planner():
    wb = Workbook(); ws = wb.active; ws.title = "Meals"
    sheet_title(ws, "Weekly Meal Planner", "Plan the week, then build your grocery list on the right.", 8)
    header_row(ws, 5, ["Day", "Breakfast", "Lunch", "Dinner", "Snacks"])
    for i, d in enumerate(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]):
        ws.cell(row=6 + i, column=1, value=d).font = Font(name="Lato", bold=True)
        ws.row_dimensions[6 + i].height = 32
    body(ws, 6, 12, 1, 5)
    header_row(ws, 5, ["Grocery list", "✓"], start_col=7)
    body(ws, 6, 25, 7, 8)
    widths(ws, [14, 22, 22, 22, 18, 3, 26, 5])
    return save_wb(wb, "Meal planner")


SHEETS = [monthly_budget, study_schedule, grade_tracker, chore_chart, meal_planner]


# ── Presentations ─────────────────────────────────────────────────────────────
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt as PPt


def deck():
    p = Presentation()
    p.slide_width, p.slide_height = Emu(12192000), Emu(6858000)     # 16:9
    return p


def band(slide, p, color, top=0, height=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, p.slide_width, height or p.slide_height)
    s.fill.solid()
    s.fill.fore_color.rgb = PRGB.from_string(color)
    s.line.fill.background()
    return s


def text(slide, x, y, w, h, value, size, bold=False, color="1D3557", font="Montserrat"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = par.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = PPt(size), bold, font
        r.font.color.rgb = PRGB.from_string(color)
        par.space_after = PPt(8)
    return tb


def title_slide(p, title, sub):
    s = p.slides.add_slide(p.slide_layouts[6])
    band(s, p, NAVY_HEX)
    band(s, p, BLUE_HEX, top=Emu(4600000), height=Emu(90000))
    text(s, Emu(700000), Emu(2000000), Emu(10800000), Emu(1400000), title, 48, True, "FFFFFF")
    text(s, Emu(700000), Emu(3500000), Emu(10800000), Emu(800000), sub, 20, False, "BFD7F2", "Lato")
    return s


def content_slide(p, title, bullets):
    s = p.slides.add_slide(p.slide_layouts[6])
    band(s, p, NAVY_HEX, height=Emu(1150000))
    text(s, Emu(600000), Emu(280000), Emu(11000000), Emu(700000), title, 30, True, "FFFFFF")
    text(s, Emu(700000), Emu(1500000), Emu(10800000), Emu(4800000), ["•  " + b for b in bullets], 20, False, "1D3557", "Lato")
    return s


def save_deck(p, name):
    path = OUT / "Presentation" / (name + ".pptx")
    path.parent.mkdir(parents=True, exist_ok=True)
    p.core_properties.title = name
    p.core_properties.author = ""
    p.save(path)
    return path


def project_status():
    p = deck()
    title_slide(p, "[Project Name] Status Update", "[Team] · [Date]")
    content_slide(p, "Summary", ["Overall status: [On track / At risk / Off track]", "[One-sentence summary of progress]", "[Key decision needed, if any]"])
    content_slide(p, "Progress since last update", ["[Milestone reached]", "[Work completed]", "[Metric: X of Y done]"])
    content_slide(p, "Risks and issues", ["[Risk] – [impact] – [mitigation]", "[Issue] – [owner] – [due date]"])
    content_slide(p, "Next steps", ["[Next milestone and date]", "[Upcoming work]", "[Help needed from stakeholders]"])
    return save_deck(p, "Project status")


def class_presentation():
    p = deck()
    title_slide(p, "[Topic of Your Presentation]", "[Your Name] · [Class] · [Date]")
    content_slide(p, "What I'll cover", ["[Question or problem]", "[Background]", "[What I found]", "[Conclusion]"])
    content_slide(p, "Background", ["[Key fact or context]", "[Why this topic matters]"])
    content_slide(p, "What I found", ["[Finding one, with evidence]", "[Finding two, with evidence]", "[Finding three, with evidence]"])
    content_slide(p, "Conclusion", ["[Answer to the question]", "[What I learned]", "[Question for the audience]"])
    content_slide(p, "Sources", ["[Author, Title, Year]", "[Website name, page title, date accessed]"])
    return save_deck(p, "Class presentation")


def photo_album():
    p = deck()
    title_slide(p, "[Album Title]", "[Occasion] · [Date or place]")
    for caption in ("[Caption for this photo]", "[Caption for this photo]", "[Caption for this photo]"):
        s = p.slides.add_slide(p.slide_layouts[6])
        band(s, p, "F4F6FA")
        frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(1400000), Emu(500000), Emu(9400000), Emu(5000000))
        frame.fill.solid()
        frame.fill.fore_color.rgb = PRGB.from_string("DDE6F1")
        frame.line.color.rgb = PRGB.from_string("FFFFFF")
        frame.line.width = Emu(76200)
        frame.text_frame.text = "Insert → Image to add a photo here"
        frame.text_frame.paragraphs[0].runs[0].font.color.rgb = PRGB.from_string("5F6B7A")
        text(s, Emu(1400000), Emu(5700000), Emu(9400000), Emu(600000), caption, 22, False, "1D3557", "Caveat")
    return save_deck(p, "Photo album")


DECKS = [project_status, class_presentation, photo_album]


def main():
    made = [f() for f in DOCUMENTS + SHEETS + DECKS]
    for m in made:
        print("  ", m.relative_to(OUT))
    print(len(made), "templates")


if __name__ == "__main__":
    main()
